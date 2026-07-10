import docx
import re
from pathlib import Path
from docx.shared import Pt, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH

# Source and destination
md_path = Path(r"C:\Users\ben.lefort\.gemini\antigravity\brain\4707f54b-be72-4525-a5a3-f7166d8b5a4d\ofa_fertilizer_policy_memo.md")
docx_path = Path(r"C:\Users\ben.lefort\OneDrive - Ontario Federation of Agriculture\Desktop\Ben Desktop Files\Economic Analyst Position\Input Cost Research\Fertilizer\OFA_Fertilizer_Policy_Memo_v7.docx")

content = md_path.read_text(encoding="utf-8")

# Initialize document
doc = docx.Document()
style = doc.styles['Normal']
font = style.font
font.name = 'Calibri'
font.size = Pt(11)

# Images to embed when certain sections are reached
img15_path = r"C:\Projects\Farm Finance Stats Dashboard\Database\farm_finance_dashboard_starter\data\latest\tariff_research\charts\15_ontario_vs_west.png"
img16_path = r"C:\Projects\Farm Finance Stats Dashboard\Database\farm_finance_dashboard_starter\data\latest\tariff_research\charts\16_regional_heatmap.png"

# Basic Markdown parsing
lines = content.split('\n')

for line in lines:
    line = line.strip()
    if not line:
        continue
    
    # HR rule
    if line.startswith('---'):
        doc.add_paragraph().add_run('_'*50)
        continue
        
    # Headers
    if line.startswith('# '):
        p = doc.add_paragraph()
        run = p.add_run(line[2:])
        run.bold = True
        run.font.size = Pt(16)
        continue
    elif line.startswith('## '):
        p = doc.add_paragraph()
        run = p.add_run(line[3:])
        run.bold = True
        run.font.size = Pt(14)
        
        # Inject visuals under specific headers
        if "The Structural Deficit:" in line:
            import os
            if os.path.exists(img15_path):
                img_p = doc.add_paragraph()
                img_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                img_p.add_run().add_picture(img15_path, width=Inches(6.0))
                
                txt_p = doc.add_paragraph()
                txt_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                txt_run = txt_p.add_run("Figure 1: Ontario vs Western Canada Nitrogen Fertilizer Price Index")
                txt_run.italic = True
                txt_run.font.size = Pt(9)
            if os.path.exists(img16_path):
                img_p2 = doc.add_paragraph()
                img_p2.alignment = WD_ALIGN_PARAGRAPH.CENTER
                img_p2.add_run().add_picture(img16_path, width=Inches(6.0))
                
                txt_p2 = doc.add_paragraph()
                txt_p2.alignment = WD_ALIGN_PARAGRAPH.CENTER
                txt_run2 = txt_p2.add_run("Figure 2: Regional Pre and Post-Tariff Heatmap")
                txt_run2.italic = True
                txt_run2.font.size = Pt(9)
        continue
        
    elif line.startswith('### '):
        p = doc.add_paragraph()
        run = p.add_run(line[4:])
        run.bold = True
        run.font.size = Pt(12)
        continue
    elif line.startswith('#### '):
        p = doc.add_paragraph()
        run = p.add_run(line[5:])
        run.bold = True
        run.italic = True
        continue
        
    # Lists
    if line.startswith('* ') or line.startswith('- '):
        p = doc.add_paragraph(style='List Bullet')
        text = line[2:]
    elif re.match(r'^\d+\.\s', line):
        p = doc.add_paragraph(style='List Number')
        text = re.sub(r'^\d+\.\s', '', line)
    else:
        p = doc.add_paragraph()
        text = line
        
    # Process inline formatting (bold)
    parts = re.split(r'(\*\*.*?\*\*)', text)
    for part in parts:
        if part.startswith('**') and part.endswith('**'):
            p.add_run(part[2:-2]).bold = True
        elif part.startswith('*') and part.endswith('*'):
            p.add_run(part[1:-1]).italic = True
        elif part.startswith('_') and part.endswith('_'):
            p.add_run(part[1:-1]).italic = True
        else:
            if part.startswith('**TO:**') or part.startswith('**FROM:**') or part.startswith('**DATE:**') or part.startswith('**SUBJECT:**'):
                p.add_run(part).bold = False
                p.text = ""
                key_match = re.match(r'\*\*(.*?)\*\*\s*(.*)', part)
                if key_match:
                    r1 = p.add_run(key_match.group(1) + " ")
                    r1.bold = True
                    r2 = p.add_run(key_match.group(2))
            else:
                p.add_run(part)

doc.save(docx_path)
print(f"Successfully saved to {docx_path}")
