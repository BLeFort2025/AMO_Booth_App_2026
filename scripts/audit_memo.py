"""
Full document audit — scan every paragraph and table cell for dollar amounts
and percentage claims, flag any that may need updating.
"""

import re
from docx import Document

DOC_PATH = (
    r"C:\Users\ben.lefort\OneDrive - Ontario Federation of Agriculture"
    r"\Desktop\Ben Desktop Files\Farm Business Confidence Survey"
    r"\2025\Specififc Issue spekaing notes"
    r"\Energy and Fertilizer Shock Memo and Speaking Notes.docx"
)

doc = Document(DOC_PATH)

# Known correct values (post MAP-N fix)
CORRECT = {
    # Scenario 1
    "grain_2025_total": 97563,
    "grain_2026_cons": 127993,
    "grain_2026_mid": 136685,
    "grain_2026_worst": 145692,
    "grain_increase_cons": 30430,
    "grain_increase_mid": 39121,
    "grain_increase_worst": 48128,
    "grain_per_acre_cons": 61,
    "grain_per_acre_mid": 78,
    "grain_per_acre_worst": 96,
    "grain_fert_2025": 65182,
    "grain_fert_2026_mid": 84420,
    # Scenario 2 (unchanged)
    "veg_2025": 21825,
    "veg_2026_cons": 31659,
    "veg_2026_mid": 32325,
    "veg_2026_worst": 33551,
}

# OLD incorrect values to flag
OLD_WRONG = {
    "$101,022": "OLD grain 2025 total — should be $97,563",
    "$101,000": "OLD grain 2025 total (rounded) — should be ~$98,000",
    "$141,808": "OLD grain 2026 mid — should be $136,685",
    "$151,010": "OLD grain 2026 worst — should be $145,692",
    "$133,011": "OLD grain 2026 cons — should be $127,993",
    "$40,785": "OLD grain increase mid — should be $39,121",
    "$49,987": "OLD grain increase worst — should be $48,128",
    "$31,989": "OLD grain increase cons — should be $30,430",
    "$68,641": "OLD grain fert 2025 — should be $65,182",
    "$89,543": "OLD grain fert 2026 mid — should be $84,420",
    "$82/ac": "OLD grain per-acre mid — should be $78/ac",
    "$64/ac": "OLD grain per-acre cons — should be $61/ac",
    "$100/ac": "OLD grain per-acre worst — should be $96/ac",
    "$82 per acre": "OLD grain per-acre mid — should be $78 per acre",
    "+32%": "OLD grain increase % cons — should be +31%",
}

# Scan all paragraphs
print("=" * 80)
print("FULL DOCUMENT AUDIT")
print("=" * 80)

issues = []
all_lines = []

print("\n--- PARAGRAPHS ---")
for i, para in enumerate(doc.paragraphs):
    text = para.text.strip()
    if not text:
        continue
    
    all_lines.append(("PARA", i, text))
    
    # Check for old wrong values
    for old_val, description in OLD_WRONG.items():
        if old_val in text:
            issues.append(f"PARA {i}: STALE VALUE '{old_val}' — {description}")
            issues.append(f"  Text: {text[:150]}...")

    # Check for any dollar amounts to review
    dollar_matches = re.findall(r'\$[\d,]+(?:\.\d+)?(?:/\w+)?', text)
    pct_matches = re.findall(r'\d+(?:\.\d+)?%?\s*(?:percent|%)', text)

print("\n--- TABLE CELLS ---")
for t_idx, table in enumerate(doc.tables):
    for r_idx, row in enumerate(table.rows):
        for c_idx, cell in enumerate(row.cells):
            text = cell.text.strip()
            if not text:
                continue
            
            all_lines.append(("TABLE", f"{t_idx}.{r_idx}.{c_idx}", text))
            
            for old_val, description in OLD_WRONG.items():
                if old_val in text:
                    issues.append(f"TABLE {t_idx} R{r_idx} C{c_idx}: STALE VALUE '{old_val}' — {description}")
                    issues.append(f"  Text: {text[:150]}")

# Print all content with dollar/pct values
print("\n" + "=" * 80)
print("ALL LINES WITH DOLLAR AMOUNTS OR PERCENTAGES")
print("=" * 80)
for loc_type, loc_id, text in all_lines:
    has_dollar = "$" in text
    has_pct = "%" in text or "percent" in text.lower()
    if has_dollar or has_pct:
        print(f"\n[{loc_type} {loc_id}]")
        print(f"  {text[:200]}")

# Print issues
print("\n" + "=" * 80)
if issues:
    print(f"ISSUES FOUND: {len(issues)}")
    print("=" * 80)
    for issue in issues:
        print(f"  {issue}")
else:
    print("NO STALE VALUES FOUND — All numbers appear current")
    print("=" * 80)
