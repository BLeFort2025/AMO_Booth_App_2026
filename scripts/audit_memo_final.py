"""Post-OMAFRA rebuild audit."""
from docx import Document

DOC_PATH = (
    r"C:\Users\ben.lefort\OneDrive - Ontario Federation of Agriculture"
    r"\Desktop\Ben Desktop Files\Farm Business Confidence Survey"
    r"\2025\Specififc Issue spekaing notes"
    r"\Energy and Fertilizer Shock Memo and Speaking Notes.docx"
)

doc = Document(DOC_PATH)

# Values that should NOT be in the document anymore
STALE = {
    "$97,563": "OLD rotation total",
    "$101,022": "VERY OLD rotation total",
    "$165,740": "OLD corn-25L total",
    "$164,522": "OLD corn-22.9L total",
    "$136,685": "OLD rotation 2026 mid",
    "$234,865": "OLD corn-25L 2026 mid",
    "$232,608": "OLD corn-22.9L 2026 mid",
    "+40%": "OLD pct (may be valid in topline)",
    "+42%": "OLD corn-25L pct",
    "$166,000": "OLD corn-25L rounded",
    "$220,000": "OLD corn-25L cons rounded",
    "$251,000": "OLD corn-25L worst rounded",
    "corn, soybean": "OLD rotation",
    "200/175/125": "OLD rotation split",
    "minimum-till": "OLD tillage (now no-till)",
    "min-till": "OLD tillage",
}

issues = []

print("=" * 80)
print("ALL $ AND % VALUES IN DOCUMENT")
print("=" * 80)

for i, para in enumerate(doc.paragraphs):
    text = para.text.strip()
    if not text:
        continue
    if "$" in text or "percent" in text.lower() or "%" in text:
        print(f"\n[P{i}] {text[:200]}")
    for s, d in STALE.items():
        if s == "+40%" and ("40 percent" in text or "over 40" in text):
            continue  # speaking notes say "over 40 percent" which is still broadly correct
        if s in text:
            issues.append(f"  PARA {i}: STALE '{s}' ({d}) in: {text[:120]}")

for t_idx, table in enumerate(doc.tables):
    for r_idx, row in enumerate(table.rows):
        for c_idx, cell in enumerate(row.cells):
            text = cell.text.strip()
            if "$" in text or "%" in text:
                print(f"[T{t_idx}.R{r_idx}.C{c_idx}] {text}")
            for s, d in STALE.items():
                if s in text:
                    issues.append(f"  TABLE {t_idx} R{r_idx} C{c_idx}: STALE '{s}' ({d}) -> {text}")

print("\n" + "=" * 80)
if issues:
    print(f"ISSUES ({len(issues)}):")
    for i in issues:
        print(i)
else:
    print("ALL CLEAR")
print("=" * 80)
