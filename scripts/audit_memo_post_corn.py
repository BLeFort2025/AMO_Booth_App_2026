"""
Post-corn-update audit - check every line for stale rotation-era values.
"""
import re
from docx import Document

DOC_PATH = (
    r"C:\Users\ben.lefort\OneDrive - Ontario Federation of Agriculture"
    r"\Desktop\Ben Desktop Files\Farm Business Confidence Survey"
    r"\2025\Specififc Issue spekaing notes"
    r"\Energy and Fertilizer Shock Memo and Speaking Notes.docx"
)

doc = Document(DOC_PATH)

# STALE values from the old rotation model that should NOT appear
STALE = {
    "$97,563": "OLD rotation total 2025",
    "$98,000": "OLD rotation total 2025 rounded",
    "$101,022": "VERY OLD rotation total 2025",
    "$101,000": "VERY OLD rotation total 2025 rounded",
    "$136,685": "OLD rotation 2026 mid",
    "$128,000": "OLD rotation 2026 cons rounded",
    "$146,000": "OLD rotation 2026 worst rounded",
    "$39,121": "OLD rotation increase mid",
    "$30,000 to $48,000": "OLD rotation increase range",
    "$30,430": "OLD rotation increase cons",
    "$48,128": "OLD rotation increase worst",
    "$68,641": "VERY OLD fert 2025",
    "$65,182": "OLD rotation fert 2025",
    "$78 per acre": "OLD rotation per-acre mid",
    "corn, soybean, and wheat": "OLD rotation description",
    "corn-soybean-wheat": "OLD rotation description",
    "Corn-soybean-wheat": "OLD rotation description",
    "200/175/125": "OLD rotation split",
    "soybean": "OLD rotation crop mention in grain scenario context",
}

issues = []

# Check paragraphs
for i, para in enumerate(doc.paragraphs):
    text = para.text.strip()
    if not text:
        continue
    for stale_val, desc in STALE.items():
        # Skip "soybean" check in Q&A about switching to soybeans
        if stale_val == "soybean" and ("Q:" in text or "switch to" in text.lower() or "fix their own" in text.lower()):
            continue
        if stale_val in text:
            issues.append(f"PARA {i}: Found '{stale_val}' ({desc})")
            issues.append(f"  -> {text[:180]}")

# Check table cells
for t_idx, table in enumerate(doc.tables):
    for r_idx, row in enumerate(table.rows):
        for c_idx, cell in enumerate(row.cells):
            text = cell.text.strip()
            for stale_val, desc in STALE.items():
                if stale_val == "soybean":
                    continue
                if stale_val in text:
                    issues.append(f"TABLE {t_idx} R{r_idx} C{c_idx}: Found '{stale_val}' ({desc})")
                    issues.append(f"  -> {text[:120]}")

# Print ALL paragraphs with dollar amounts for manual review
print("=" * 80)
print("ALL PARAGRAPHS WITH DOLLAR AMOUNTS (post-update)")
print("=" * 80)
for i, para in enumerate(doc.paragraphs):
    text = para.text.strip()
    if "$" in text:
        print(f"\n[PARA {i}] {text[:220]}")

print("\n" + "=" * 80)
if issues:
    print(f"STALE VALUES FOUND: {len(issues)}")
    print("=" * 80)
    for issue in issues:
        print(f"  {issue}")
else:
    print("ALL CLEAR - No stale values detected")
    print("=" * 80)
