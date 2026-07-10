"""
Generate Health Access Developer Manual (Word .docx)
====================================================
Produces a comprehensive developer manual covering the Health Access
section's data pipeline, UI rendering, and maintenance procedures.

Run:  python scripts/generate_health_manual.py
"""

from docx import Document
from docx.shared import Pt
from docx.enum.table import WD_TABLE_ALIGNMENT
from pathlib import Path

OUTPUT = Path(
    r"C:\Projects\Farm Finance Stats Dashboard"
    r"\Reports and supplemental information"
    r"\Rural Community Data Tabs"
    r"\Health Access Tab - Developer Manual.docx"
)


def _add_heading(doc, text, level=1):
    return doc.add_heading(text, level=level)


def _add_para(doc, text, bold=False, italic=False, style=None):
    p = doc.add_paragraph(style=style)
    run = p.add_run(text)
    run.bold = bold
    run.italic = italic
    return p


def _add_table(doc, headers, rows):
    t = doc.add_table(rows=1 + len(rows), cols=len(headers))
    t.style = "Light Grid Accent 1"
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    for i, h in enumerate(headers):
        cell = t.rows[0].cells[i]
        cell.text = h
        for p in cell.paragraphs:
            for r in p.runs:
                r.bold = True
    for ri, row in enumerate(rows):
        for ci, val in enumerate(row):
            t.rows[ri + 1].cells[ci].text = str(val)
    return t


def _add_code(doc, code_text):
    p = doc.add_paragraph()
    run = p.add_run(code_text)
    run.font.name = "Consolas"
    run.font.size = Pt(9)
    run.font.color.rgb = None
    return p


def build_manual():
    doc = Document()
    doc.add_heading("🏥 Health Access Tab — Developer Manual", level=0)

    doc.add_paragraph(
        "This manual documents the complete architecture, data pipeline, UI rendering, "
        "and maintenance procedures for the Health Access section (§9) of the "
        "Rural Community Data page."
    )
    _add_para(doc,
        "Document Version: 1.0  |  Date: February 2026  |  "
        "Status: APPROVED FOR PRODUCTION (Red Team Round 3 Sign-Off)",
        bold=True,
    )

    # ══════════════════════════════════════════════════════════════════════
    # 1. OVERVIEW
    # ══════════════════════════════════════════════════════════════════════
    _add_heading(doc, "1. Overview", level=1)
    doc.add_paragraph(
        "The Health Access section provides per-community health facility counts "
        "for Ontario Census Subdivisions (CSDs). It displays facility breakdowns "
        "by type (Hospitals, Ambulatory Care, Nursing & Residential, Other), "
        "per-capita rates, and provincial/rural benchmarks."
    )
    doc.add_paragraph(
        "Data source: Statistics Canada Open Database of Healthcare Facilities (ODHF), "
        "a point-in-time open-data release cataloguing physical healthcare facility "
        "locations across Canada."
    )

    # ══════════════════════════════════════════════════════════════════════
    # 2. FILE INVENTORY
    # ══════════════════════════════════════════════════════════════════════
    _add_heading(doc, "2. File Inventory", level=1)
    _add_table(doc,
        ["File", "Purpose", "Type"],
        [
            ["scripts/fetch_health_facilities.py", "ETL script — processes ODHF ZIP", "Manual pipeline"],
            ["app/pages/9_🏘️_Rural_Community_Data.py", "Streamlit UI (§9 Health Access, lines 2319–2480)", "Frontend"],
            ["data/raw/ODHF_v1.1.zip", "Raw source data (manual download)", "Input"],
            ["data/latest/wellbeing/health_facility_counts.csv", "WIDE output (one row per CSD)", "Output"],
            ["data/derived/health_facilities.csv", "LONG output (sgc_code/indicator/value)", "Output"],
            ["data/latest/wellbeing/amenities_vintage.json", "Runtime vintage metadata", "Metadata"],
            ["data/latest/wellbeing/dim_geography.csv", "Community names and counties", "Dependency"],
            ["data/latest/wellbeing/census_indicators.csv", "Population data for per-capita rates", "Dependency"],
        ],
    )

    # ══════════════════════════════════════════════════════════════════════
    # 3. ARCHITECTURE
    # ══════════════════════════════════════════════════════════════════════
    _add_heading(doc, "3. Architecture", level=1)
    doc.add_paragraph(
        "The Health Access pipeline follows a manual-trigger, static-source model. "
        "Unlike automated pipelines (Census, Broadband), this pipeline requires a "
        "one-time manual download of the ODHF ZIP file from Statistics Canada's LODE "
        "page, followed by a manual script invocation."
    )
    _add_heading(doc, "3.1 Data Flow", level=2)
    _add_code(doc,
        "ODHF_v1.1.zip (manual download from StatCan LODE)\n"
        "  → fetch_health_facilities.py (manual: python scripts/fetch_health_facilities.py)\n"
        "  → WIDE CSV: data/latest/wellbeing/health_facility_counts.csv\n"
        "       One row per Ontario CSD (including zero-facility communities)\n"
        "  → LONG CSV: data/derived/health_facilities.csv\n"
        "       Legacy sgc_code/indicator/value format\n"
        "  → META JSON: data/latest/wellbeing/amenities_vintage.json\n"
        "       Dynamic vintage metadata read by UI at runtime\n"
        "\n"
        "Consumer: Standalone '🏥 Health Access' section (§9)\n"
        "INDICATOR_META: health_facilities (quarantined: standalone-only view)"
    )

    _add_heading(doc, "3.2 Why Not Automated?", level=2)
    doc.add_paragraph(
        "The ODHF is released as a static ZIP file on Statistics Canada's Linked Open "
        "Data Environment (LODE). The download URL changes between releases, and there "
        "is no stable API endpoint. The data is a point-in-time snapshot (released ~2021 "
        "from ~2019-2020 collection) and does not update frequently. Automating download "
        "would require URL maintenance with no real benefit given the data's static nature."
    )

    _add_heading(doc, "3.3 INDICATOR_META Quarantine", level=2)
    doc.add_paragraph(
        "The health_facilities indicator is registered in INDICATOR_META with the "
        "metadata flag {\"supported_views\": [\"standalone\"]}. This quarantine prevents "
        "the indicator from appearing in generic time-series tabs (which expect Census "
        "data with census_year columns). Health data is a point-in-time snapshot, not a "
        "time-series — it only renders in the standalone §9 section."
    )

    # ══════════════════════════════════════════════════════════════════════
    # 4. ETL PIPELINE DETAILS
    # ══════════════════════════════════════════════════════════════════════
    _add_heading(doc, "4. ETL Pipeline Details", level=1)

    _add_heading(doc, "4.1 Input: ODHF ZIP", level=2)
    doc.add_paragraph(
        "Expected at: data/raw/ODHF_v1.1.zip\n"
        "Download from: https://www.statcan.gc.ca/en/lode/databases/odhf\n\n"
        "The script auto-detects the largest CSV inside the ZIP (by file size) and "
        "reads it with latin-1 encoding. If the ZIP is missing, the script prints a "
        "warning and exits gracefully without crashing the pipeline."
    )

    _add_heading(doc, "4.2 Column Resolution (Alias System)", level=2)
    doc.add_paragraph(
        "StatCan changes column header casing between ODHF releases (CSDuid → CSDUID "
        "→ csd_uid). To handle this, the script normalizes all columns to lowercase "
        "and resolves against explicit alias lists:"
    )
    _add_table(doc,
        ["Column", "Aliases", "Purpose"],
        [
            ["CSD UID", "csduid, csd_uid, csd", "Census Subdivision identifier"],
            ["Province", "province, prov_terr, province / territory", "Ontario filter"],
            ["Facility Type", "odhf_facility_type, facility_type, source_facility_type", "Type classification"],
        ],
    )
    doc.add_paragraph(
        "If no alias matches, the script raises a RuntimeError with all available "
        "column names, enabling immediate debugging."
    )

    _add_heading(doc, "4.3 Ontario Filtering", level=2)
    doc.add_paragraph(
        "The province column is matched against ['on', 'ontario', '35'] (case-insensitive) "
        "to capture all possible ODHF province encodings. CSD codes are then validated "
        "against the regex ^35\\d{5}$ to ensure exactly 7-digit Ontario SGC codes."
    )

    _add_heading(doc, "4.4 Facility Type Mapping", level=2)
    _add_table(doc,
        ["ODHF Type (lowercase)", "Internal Column", "Description"],
        [
            ["hospitals", "hospitals", "Acute care hospitals"],
            ["ambulatory health care services", "ambulatory", "Walk-in clinics, family health teams"],
            ["nursing and residential care facilities", "nursing_residential", "Long-term care, retirement homes"],
            ["(all other values)", "other", "Catch-all for unmapped types"],
        ],
    )
    doc.add_paragraph(
        "The 'other' bin captures any ODHF facility types not in the explicit map. "
        "When 'other' facilities are detected, the script logs the unmapped type names "
        "and counts for debugging."
    )

    _add_heading(doc, "4.5 Outer Join: Zero-Facility CSDs", level=2)
    doc.add_paragraph(
        "CRITICAL DESIGN DECISION (P0 Audit Fix): The ETL performs an OUTER join "
        "with dim_geography.csv to include ALL Ontario CSDs in the output, even those "
        "with zero health facilities. Without this, the 'Rural Ontario Average' "
        "benchmark was subject to survivorship bias — only counting the populations "
        "of municipalities that already had facilities, which inflated the per-capita "
        "rate."
    )
    _add_code(doc,
        "# Outer join ensures zero-facility CSDs appear with explicit 0s\n"
        "counts = counts.merge(\n"
        "    geo_on[['sgc_code', 'geo_name', 'county']],\n"
        "    on='sgc_code', how='outer',\n"
        ")\n"
        "# Backfill NaN facility counts with 0 for newly-joined CSDs\n"
        "for col in cat_cols:\n"
        "    counts[col] = counts[col].fillna(0).astype(int)"
    )
    doc.add_paragraph(
        "This increases the CSV from ~150 rows (facility-bearing CSDs only) to ~444+ rows "
        "(all Ontario CSDs), providing accurate denominators for per-capita calculations."
    )

    _add_heading(doc, "4.6 Per-Capita Rate Calculation", level=2)
    doc.add_paragraph(
        "The ETL merges population data from census_indicators.csv (latest Census year) "
        "and computes facilities_per_10k = (total_facilities / population) × 10,000. "
        "Edge cases are handled explicitly:"
    )
    _add_table(doc,
        ["Edge Case", "Handling", "Result"],
        [
            ["population = 0", "Division produces inf", "Replaced with None"],
            ["population = NaN", "Division produces NaN", "Replaced with None (np.nan added to replace list)"],
            ["population missing entirely", "census_indicators.csv absent", "Column not created; UI skips per-capita mode"],
        ],
    )

    _add_heading(doc, "4.7 Community Name Generation", level=2)
    doc.add_paragraph(
        "Community names are generated as 'geo_name (county)' format. The lambda "
        "checks BOTH geo_name AND county are non-NaN before concatenating, preventing "
        "literal 'nan' strings (I14 audit fix). Fallback cascade:"
    )
    _add_code(doc,
        "# Priority: 'Township Name (County)' → 'Township Name' → '3523043'\n"
        "if pd.notna(geo_name) and pd.notna(county) and county:\n"
        "    community = f'{geo_name} ({county})'\n"
        "elif pd.notna(geo_name):\n"
        "    community = geo_name\n"
        "else:\n"
        "    community = str(sgc_code)"
    )

    _add_heading(doc, "4.8 Dual Output", level=2)
    doc.add_paragraph(
        "WIDE format (health_facility_counts.csv): One row per CSD with columns for "
        "each facility type, total, community name, county, population, and per-capita rate. "
        "Used by the standalone §9 Health Access section."
    )
    doc.add_paragraph(
        "LONG format (health_facilities.csv): Three columns (sgc_code, indicator, value) "
        "with indicator='health_facilities' and value=total_facilities. Legacy format "
        "for potential use in generic indicator views (currently quarantined)."
    )

    _add_heading(doc, "4.9 Vintage Metadata", level=2)
    doc.add_paragraph(
        "The ETL writes/updates amenities_vintage.json with ODHF metadata including "
        "source name, version, release year, pipeline run timestamp, and ZIP filename. "
        "The UI reads this at runtime to display dynamic data footnotes. If the file "
        "doesn't exist, hardcoded defaults are used."
    )

    # ══════════════════════════════════════════════════════════════════════
    # 5. WIDE CSV SCHEMA
    # ══════════════════════════════════════════════════════════════════════
    _add_heading(doc, "5. Wide CSV Schema", level=1)
    _add_table(doc,
        ["Column", "Type", "Description", "Example"],
        [
            ["sgc_code", "str (7-digit)", "Census Subdivision identifier", "3523043"],
            ["ambulatory", "int", "Count of ambulatory care facilities", "6"],
            ["hospitals", "int", "Count of hospitals", "2"],
            ["nursing_residential", "int", "Count of nursing/residential facilities", "13"],
            ["other", "int", "Count of unmapped facility types", "0"],
            ["total_facilities", "int", "Sum of all facility types", "21"],
            ["geo_name", "str", "Community name from dim_geography.csv", "Cornwall"],
            ["county", "str", "County/district name", "Stormont, Dundas and Glengarry"],
            ["community", "str", "Formatted display name", "Cornwall (Stormont, Dundas and Glengarry)"],
            ["population", "float", "Latest Census population (may be NaN)", "47845.0"],
            ["facilities_per_10k", "float/None", "Per 10,000 population rate", "4.4"],
        ],
    )

    # ══════════════════════════════════════════════════════════════════════
    # 6. UI RENDERING
    # ══════════════════════════════════════════════════════════════════════
    _add_heading(doc, "6. UI Rendering", level=1)

    _add_heading(doc, "6.1 Data Loading", level=2)
    doc.add_paragraph(
        "load_health_data() reads health_facility_counts.csv with @st.cache_data(ttl=3600). "
        "SGC codes are re-normalized (pd.to_numeric → int → str.zfill(7)) and filtered "
        "to Ontario (startswith '35'). Returns None if the file doesn't exist."
    )

    _add_heading(doc, "6.2 Layout", level=2)
    doc.add_paragraph(
        "The section uses a 3:2 column layout:"
    )
    _add_table(doc,
        ["Column", "Content"],
        [
            ["Left (60%)", "Grouped bar chart with metric toggle and benchmark lines"],
            ["Right (40%)", "Summary data table with Ontario average caption"],
        ],
    )

    _add_heading(doc, "6.3 Metric Toggle", level=2)
    doc.add_paragraph(
        "A horizontal radio widget lets users switch between 'Per 10K Population' and "
        "'Raw Counts'. The toggle only appears if the facilities_per_10k column exists "
        "in the data. The selected mode controls: (a) which values are plotted, "
        "(b) whether benchmark lines appear, (c) whether asterisk footnotes render."
    )

    _add_heading(doc, "6.4 Chart Details", level=2)
    _add_table(doc,
        ["Element", "Detail"],
        [
            ["Chart type", "px.bar with barmode='group'"],
            ["Facility types", "Hospitals, Ambulatory, Nursing & Residential, Other"],
            ["Colors", "#e74c3c (red), #3498db (blue), #27ae60 (green), #f39c12 (orange)"],
            ["Template", "plotly_white"],
            ["Height", "400px"],
            ["Legend", "Horizontal, bottom-centered"],
        ],
    )

    _add_heading(doc, "6.5 Benchmark Lines", level=2)
    doc.add_paragraph(
        "Two benchmark lines appear in per-capita mode only:"
    )
    _add_table(doc,
        ["Line", "Style", "Color", "Formula"],
        [
            ["Ontario Average", "Dashed", "#9e9e9e (gray)",
             "valid_pop_df['total_facilities'].sum() / valid_pop_df['population'].sum() × 10,000"],
            ["Rural Ontario (<30K)", "Dotted", "#66bb6a (green)",
             "rural['total_facilities'].sum() / rural['population'].sum() × 10,000 (pop < 30K)"],
        ],
    )
    doc.add_paragraph(
        "CRITICAL: The Ontario average uses dropna(subset=['population']) to ensure "
        "strict numerator/denominator alignment. CSDs with unknown populations are "
        "excluded from BOTH the facility count and the population total, preventing "
        "artificial inflation of the rate (Q9 audit fix)."
    )

    _add_heading(doc, "6.6 NaN Population Handling (Q7)", level=2)
    doc.add_paragraph(
        "In per-capita mode, CSDs with NaN or zero population produce display_val=None. "
        "Plotly gracefully omits these bars rather than plotting raw facility counts on "
        "a per-capita axis, which would massively distort the scale. In raw counts mode, "
        "all CSDs are displayed regardless of population data availability."
    )

    _add_heading(doc, "6.7 Small-Population Asterisk", level=2)
    doc.add_paragraph(
        "Communities with population < 1,000 receive an asterisk suffix (*) on per-capita "
        "charts only. The corresponding footnote ('Per-capita rates for these communities "
        "are statistically volatile...') renders conditionally — only when the per-capita "
        "toggle is active (Q12 audit fix)."
    )

    _add_heading(doc, "6.8 Community Names (Q1, Q10)", level=2)
    doc.add_paragraph(
        "Full community names (including county) are retained in both chart and table — "
        "no truncation via .split('(') applied. This prevents ambiguity for Ontario "
        "municipalities that share names (e.g., Wellesley township vs. settlement). "
        "The UI prefers ETL-generated community names but falls back to community_lookup "
        "if the column is missing (Q10 resiliency fix)."
    )

    _add_heading(doc, "6.9 Captions and Footnotes", level=2)
    _add_table(doc,
        ["Caption", "Content"],
        [
            ["Vintage", "Source name, version, release year from amenities_vintage.json"],
            ["Staleness", "'Facility operational status may have changed since release year' (Q8)"],
            ["Boundaries", "'Facility counts based on municipal boundaries; zero may indicate regional hub access'"],
            ["Asterisk", "Conditional: small-population statistical volatility warning (per-capita mode only)"],
        ],
    )

    # ══════════════════════════════════════════════════════════════════════
    # 7. HELPER FUNCTIONS
    # ══════════════════════════════════════════════════════════════════════
    _add_heading(doc, "7. Helper Functions", level=1)
    _add_table(doc,
        ["Function", "Location", "Cache", "Purpose"],
        [
            ["load_health_data()", "Lines 243-254", "ttl=3600", "Load & normalize health CSV"],
            ["_get_ontario_population()", "Lines 533-552", "ttl=600",
             "Total Ontario pop from Census (used by Arts section; Health uses dropna path)"],
            ["_get_rural_totals(df)", "Lines 555-563", "None",
             "Sum rural pop + facilities for CSDs < 30K"],
            ["_get_amenity_vintage(section)", "Lines 511-530", "ttl=600",
             "Read dynamic vintage metadata from JSON"],
        ],
    )

    # ══════════════════════════════════════════════════════════════════════
    # 8. DATA DEPENDENCIES
    # ══════════════════════════════════════════════════════════════════════
    _add_heading(doc, "8. Data Dependencies", level=1)
    _add_table(doc,
        ["Dependency", "Required?", "Effect if Missing"],
        [
            ["ODHF_v1.1.zip", "Yes", "ETL prints warning, exits; no CSV produced"],
            ["dim_geography.csv", "Recommended",
             "Community names fallback to SGC codes; UI uses community_lookup"],
            ["census_indicators.csv", "Recommended",
             "No population data; per-capita mode unavailable; Ontario avg hidden"],
        ],
    )

    # ══════════════════════════════════════════════════════════════════════
    # 9. MAINTENANCE PROCEDURES
    # ══════════════════════════════════════════════════════════════════════
    _add_heading(doc, "9. Maintenance Procedures", level=1)

    _add_heading(doc, "9.1 Updating ODHF Data", level=2)
    doc.add_paragraph(
        "When Statistics Canada releases a new ODHF version:"
    )
    doc.add_paragraph(
        "1. Download the new ZIP from https://www.statcan.gc.ca/en/lode/databases/odhf",
        style="List Number"
    )
    doc.add_paragraph(
        "2. Save to data/raw/ (e.g., ODHF_v2.0.zip)",
        style="List Number"
    )
    doc.add_paragraph(
        "3. Update ODHF_ZIP constant in fetch_health_facilities.py to point to the new file",
        style="List Number"
    )
    doc.add_paragraph(
        "4. Check if column headers have changed — run with --dry-run or check error output. "
        "If new aliases are needed, add them to CSD_ALIASES, PROV_ALIASES, or TYPE_ALIASES",
        style="List Number"
    )
    doc.add_paragraph(
        "5. Check if new facility types exist — the 'other' bin will catch them automatically. "
        "Decide if new types warrant explicit mapping in FACILITY_TYPE_MAP",
        style="List Number"
    )
    doc.add_paragraph(
        "6. Run: python scripts/fetch_health_facilities.py",
        style="List Number"
    )
    doc.add_paragraph(
        "7. Update version in _write_vintage() to match the new ODHF release version",
        style="List Number"
    )
    doc.add_paragraph(
        "8. Verify the output CSV row count increased (all Ontario CSDs should be present)",
        style="List Number"
    )

    _add_heading(doc, "9.2 Census Year Updates", level=2)
    doc.add_paragraph(
        "Per-capita rates use the latest Census year from census_indicators.csv. When "
        "a new Census is loaded (e.g., 2026), re-running fetch_health_facilities.py will "
        "automatically pick up the newest census_year via .max(). No code changes needed."
    )

    _add_heading(doc, "9.3 Troubleshooting", level=2)
    _add_table(doc,
        ["Symptom", "Likely Cause", "Fix"],
        [
            ["Section doesn't appear", "health_facility_counts.csv missing",
             "Run the ETL script; check ODHF ZIP exists"],
            ["No per-capita toggle", "census_indicators.csv missing or no population data",
             "Ensure Census pipeline has run; check for 'population' indicator rows"],
            ["'Cannot find CSD UID column'", "ODHF headers changed in new version",
             "Add the new header to CSD_ALIASES in fetch_health_facilities.py"],
            ["Large 'other' count", "New facility type in ODHF not mapped",
             "Check unmapped type names in script output; decide if they warrant explicit mapping"],
            ["Ontario average seems high", "NaN-population CSDs contaminating rate",
             "Confirm dropna(subset=['population']) is in the pre-compute block"],
            ["Community shows as SGC code", "Missing dim_geography.csv or new CSD not in geography file",
             "Re-download geography reference data"],
        ],
    )

    # ══════════════════════════════════════════════════════════════════════
    # 10. ODHF SOURCE DETAILS
    # ══════════════════════════════════════════════════════════════════════
    _add_heading(doc, "10. ODHF Source Details", level=1)
    _add_table(doc,
        ["Property", "Value"],
        [
            ["Full Name", "Open Database of Healthcare Facilities (ODHF)"],
            ["Publisher", "Statistics Canada"],
            ["Current Version", "v1.1"],
            ["Release Year", "2021"],
            ["Data Collection Period", "~2019-2020"],
            ["Coverage", "All healthcare facilities in Canada"],
            ["Geo Resolution", "Facility point locations with CSD assignment"],
            ["License", "Open Government Licence - Canada"],
            ["Download URL", "https://www.statcan.gc.ca/en/lode/databases/odhf"],
        ],
    )

    # ══════════════════════════════════════════════════════════════════════
    # 11. EXPERT AUDIT HISTORY
    # ══════════════════════════════════════════════════════════════════════
    _add_heading(doc, "11. Expert Audit History", level=1)
    doc.add_paragraph(
        "The Health Access module underwent a comprehensive expert red team review "
        "in February 2026 across 3 rounds. The initial review returned NO-GO with "
        "2 P0 Critical, 2 P1 High, 4 P2 Medium, and 4 P3 Low findings, plus 3 "
        "additional issues discovered by the expert."
    )

    _add_heading(doc, "11.1 Round 1: Initial Review — NO-GO", level=2)
    _add_table(doc,
        ["#", "Finding", "Priority", "Status"],
        [
            ["Q7", "NaN population → raw count plotted on per-capita axis", "P0 Critical", "FIXED"],
            ["Q11", "Rural totals survivorship bias (zero-facility CSDs excluded)", "P0 Critical", "FIXED"],
            ["Q6", "Zero-facility CSDs entirely absent from dataset", "P1 High", "FIXED"],
            ["Q3", "'Other' facility type silently excluded from chart", "P1 High", "FIXED"],
            ["I13", "Silent drop of zero-facility communities in UI", "P1 High", "FIXED"],
            ["Q1", "Community name truncation creating ambiguity", "P2 Medium", "FIXED"],
            ["Q8", "Data staleness caveat insufficient", "P2 Medium", "FIXED"],
            ["Q10", "Community name override inconsistency", "P2 Medium", "FIXED"],
            ["Q12", "Unconditional asterisk footnote in raw counts mode", "P2 Medium", "FIXED"],
            ["I14", "Literal 'nan' strings in ETL community names", "P2 Medium", "FIXED"],
            ["Q2", "Rural threshold label mismatch (30K vs 100K)", "P3 Low", "FIXED"],
            ["Q4", "Pipeline not automated", "P3 Low", "ACCEPTED"],
            ["Q5", "Redundant per-capita computation", "P3 Low", "ACCEPTED"],
            ["Q9", "Ontario average computed twice", "P3 Low", "FIXED"],
            ["I15", "NaN not handled in facilities_per_10k", "P3 Low", "FIXED"],
        ],
    )

    _add_heading(doc, "11.2 Round 2: First Re-Review — Conditional GO", level=2)
    doc.add_paragraph(
        "13/15 findings returned PASS. Two FAILs identified:"
    )
    _add_table(doc,
        ["Finding", "Issue", "Fix Applied"],
        [
            ["Q9", "dropna(subset=['population']) KeyError if census_indicators.csv missing",
             "Wrapped in 'if population in health_df.columns' check"],
            ["Q12", "NameError crash if _health_metric referenced after empty selection",
             "Confirmed already initialized before conditional block"],
        ],
    )
    doc.add_paragraph(
        "Two P3 polish items addressed: Q10 resiliency fallback added; I14 lambda "
        "edge case fixed with dual pd.notna() check."
    )

    _add_heading(doc, "11.3 Round 3: Final Sign-Off — 🟢 GO", level=2)
    doc.add_paragraph(
        "All corrections verified. One final hotfix applied: Q9 Ontario average "
        "wrapped in column existence check. No regressions introduced. Streamlit "
        "radio widget state confirmed safe with pre-initialized _health_metric."
    )
    _add_para(doc,
        "Final Verdict: 🟢 APPROVED FOR PRODUCTION — SIGNED OFF (Round 3, February 2026)",
        bold=True,
    )

    # ══════════════════════════════════════════════════════════════════════
    # 12. KEY DESIGN DECISIONS
    # ══════════════════════════════════════════════════════════════════════
    _add_heading(doc, "12. Key Design Decisions", level=1)
    _add_table(doc,
        ["Decision", "Rationale"],
        [
            ["Outer join with dim_geography",
             "Includes zero-facility CSDs for accurate per-capita denominators and rural benchmarks"],
            ["Manual pipeline (not in run_pipeline.py)",
             "ODHF URLs are unstable; data is static (no automated download benefit)"],
            ["INDICATOR_META quarantine (standalone-only)",
             "Health data is point-in-time, not Census time-series; generic tabs would misrender"],
            ["Four-category facility types + 'other'",
             "Explicit ODHF mapping with catch-all prevents silent data loss"],
            ["Rural threshold = 30,000",
             "Matches Statistics Canada's small-town classification for peer-appropriate benchmarking"],
            ["dropna() for Ontario average",
             "Strict numerator/denominator alignment prevents epidemiological artifact"],
            ["NaN population → None (not raw count)",
             "Prevents massive axis distortion when plotting per-capita rates"],
            ["Preserve full community names",
             "Prevents ambiguity for Ontario municipalities sharing names"],
        ],
    )

    # ── Save ────────────────────────────────────────────────────────────
    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    doc.save(str(OUTPUT))
    print(f"✓ Manual saved to: {OUTPUT}")


if __name__ == "__main__":
    build_manual()
