"""
Ontario Farm Energy & Fertilizer Cost Impact Data Collector
============================================================
Extracts World Bank commodity price data, combines with Ridgetown 2025 baselines,
and produces a structured JSON dataset for the cost impact model.

Sources:
- World Bank CMO Pink Sheet (Monthly Prices & Indices)
- Ridgetown Farm Input Monitoring Project (April 16, 2025)
- DTN Retail Fertilizer Index (March 2026)
- NRCan / Kalibrate fuel price data
"""

import json
import os
import sys
from datetime import datetime

try:
    import openpyxl
except ImportError:
    print("Installing openpyxl...")
    os.system(f"{sys.executable} -m pip install openpyxl")
    import openpyxl


def extract_world_bank_data(filepath):
    """Extract fertilizer and energy price data from World Bank Pink Sheet."""
    wb = openpyxl.load_workbook(filepath, data_only=True)
    
    results = {
        "monthly_prices": {},
        "monthly_indices": {}
    }
    
    # --- Extract Monthly Prices ---
    ws = wb["Monthly Prices"]
    
    # Find the header row and date columns
    header_row = None
    for row in ws.iter_rows(min_row=1, max_row=10, max_col=1):
        for cell in row:
            val = str(cell.value or "").lower()
            if "commodity" in val or "1960" in val:
                header_row = cell.row
                break
        if header_row:
            break
    
    if not header_row:
        # Try row-by-row to find date pattern
        for row in ws.iter_rows(min_row=1, max_row=10):
            for cell in row:
                val = str(cell.value or "")
                if "M01" in val or "M04" in val or "2025" in val:
                    header_row = cell.row
                    break
            if header_row:
                break
    
    if not header_row:
        print("WARNING: Could not find header row in Monthly Prices sheet")
        header_row = 4  # Common default
    
    # Get date column mapping
    date_cols = {}
    target_periods = ["2025M01", "2025M02", "2025M03", "2025M04", "2025M05",
                      "2025M06", "2025M07", "2025M08", "2025M09", "2025M10",
                      "2025M11", "2025M12",
                      "2026M01", "2026M02", "2026M03", "2026M04"]
    
    for cell in ws[header_row]:
        val = str(cell.value or "")
        if val in target_periods:
            date_cols[val] = cell.column
    
    print(f"Found {len(date_cols)} date columns: {sorted(date_cols.keys())}")
    
    # Find fertilizer and energy rows
    target_commodities = {
        "urea": ["urea"],
        "dap": ["dap", "diammonium phosphate"],
        "tsp": ["tsp", "triple superphosphate"],
        "potassium_chloride": ["potassium chloride", "potash"],
        "phosphate_rock": ["phosphate rock"],
        "crude_oil_brent": ["crude oil, brent", "brent"],
        "crude_oil_wti": ["crude oil, wti", "wti"],
        "natural_gas_us": ["natural gas, us", "natural gas, henry hub"],
        "natural_gas_europe": ["natural gas, europe", "natural gas, ttf"],
    }
    
    commodity_rows = {}
    for row in ws.iter_rows(min_col=1, max_col=1, min_row=header_row + 1):
        for cell in row:
            val = str(cell.value or "").lower().strip()
            for key, patterns in target_commodities.items():
                if any(p in val for p in patterns):
                    commodity_rows[key] = cell.row
                    print(f"  Found '{key}' at row {cell.row}: {cell.value}")
    
    # Extract price data
    for commodity, row_num in commodity_rows.items():
        results["monthly_prices"][commodity] = {}
        for period, col_num in date_cols.items():
            val = ws.cell(row=row_num, column=col_num).value
            if val is not None:
                try:
                    results["monthly_prices"][commodity][period] = float(val)
                except (ValueError, TypeError):
                    results["monthly_prices"][commodity][period] = None
    
    # --- Extract Monthly Indices ---
    ws_idx = wb["Monthly Indices"]
    
    # Find header row in indices sheet
    idx_header_row = None
    for row in ws_idx.iter_rows(min_row=1, max_row=10, max_col=1):
        for cell in row:
            val = str(cell.value or "")
            if "M01" in val or "2025" in val:
                idx_header_row = cell.row
                break
        if idx_header_row:
            break
    
    if not idx_header_row:
        idx_header_row = header_row
    
    # Get date columns for indices
    idx_date_cols = {}
    for cell in ws_idx[idx_header_row]:
        val = str(cell.value or "")
        if val in target_periods:
            idx_date_cols[val] = cell.column
    
    # Find fertilizer index rows
    target_indices = {
        "fertilizers": ["fertilizers"],
        "energy": ["energy"],
        "agriculture": ["agriculture"],
    }
    
    idx_rows = {}
    for row in ws_idx.iter_rows(min_col=1, max_col=1, min_row=idx_header_row + 1):
        for cell in row:
            val = str(cell.value or "").lower().strip()
            for key, patterns in target_indices.items():
                if val in patterns:
                    idx_rows[key] = cell.row
                    print(f"  Found index '{key}' at row {cell.row}: {cell.value}")
    
    # Extract index data
    for idx_name, row_num in idx_rows.items():
        results["monthly_indices"][idx_name] = {}
        for period, col_num in idx_date_cols.items():
            val = ws_idx.cell(row=row_num, column=col_num).value
            if val is not None:
                try:
                    results["monthly_indices"][idx_name][period] = float(val)
                except (ValueError, TypeError):
                    results["monthly_indices"][idx_name][period] = None
    
    wb.close()
    return results


def build_dataset():
    """Build the complete structured dataset for the cost impact model."""
    
    # ==========================================
    # 1. BASELINE DATA (Ridgetown April 16, 2025)
    # ==========================================
    baseline_2025 = {
        "source": "University of Guelph Ridgetown Campus - Ontario Farm Input Monitoring Project",
        "survey_date": "2025-04-16",
        "currency": "CAD",
        "confidence": "HIGH",
        "url": "https://www.ridgetownc.com",
        "fertilizer": {
            "urea_46_0_0": {
                "price_per_tonne": 819.0,
                "range_low": 715.0,
                "range_high": 890.0,
                "unit": "CAD/metric tonne",
                "n_content_pct": 46.0
            },
            "map_11_52_0": {
                "price_per_tonne": 1164.0,
                "range_low": 1090.0,
                "range_high": 1205.0,
                "unit": "CAD/metric tonne",
                "n_content_pct": 11.0,
                "p2o5_content_pct": 52.0
            },
            "potash_mop_0_0_60": {
                "price_per_tonne": 678.0,
                "range_low": 615.0,
                "range_high": 740.0,
                "unit": "CAD/metric tonne",
                "k2o_content_pct": 60.0
            },
            "uan_28_0_0": {
                "price_per_tonne": 560.0,
                "range_low": 425.0,
                "range_high": 645.0,
                "unit": "CAD/metric tonne",
                "n_content_pct": 28.0
            }
        },
        "fuel": {
            "diesel_coloured_farm": {
                "price_per_litre": 1.16,
                "unit": "CAD/litre",
                "notes": "Coloured (dyed) diesel, exempt from ON provincial fuel tax"
            },
            "gasoline_regular": {
                "price_per_litre": 1.12,
                "unit": "CAD/litre"
            }
        },
        "us_comparison": {
            "source": "Ridgetown cross-border survey (MI, OH, IN)",
            "cad_usd_rate": 0.7199,
            "urea_us_avg_per_ton": 607.0,
            "map_us_avg_per_ton": 833.0,
            "potash_us_avg_per_ton": 470.0,
            "notes": "Ontario was 6-12% CHEAPER than US on currency-adjusted basis in April 2025"
        }
    }

    # ==========================================
    # 2. DTN MARCH 2026 DATA (US National Avg)
    # ==========================================
    dtn_march_2026 = {
        "source": "DTN Retail Fertilizer Index",
        "report_date": "2026-04-01",
        "period": "4th week of March 2026",
        "currency": "USD",
        "unit": "USD/US short ton (2000 lbs)",
        "confidence": "MEDIUM - US national average, not Ontario-specific",
        "url": "https://www.dtnpf.com",
        "prices": {
            "urea": {"price_per_us_ton": 826.0, "mom_change_pct": 35.0},
            "dap": {"price_per_us_ton": 857.0},
            "potash": {"price_per_us_ton": 489.0}
        }
    }

    # ==========================================
    # 3. WORLD BANK DATA (March 2026)
    # ==========================================
    wb_data_raw = {"note": "World Bank Pink Sheet not yet extracted"}
    wb_file = os.path.join(os.path.dirname(os.path.dirname(__file__)), 
                           "data", "CMO-Historical-Data-Monthly.xlsx")
    if os.path.exists(wb_file):
        print(f"\nExtracting World Bank data from: {wb_file}")
        try:
            wb_data_raw = extract_world_bank_data(wb_file)
        except Exception as e:
            print(f"Error extracting World Bank data: {e}")
            wb_data_raw = {"error": str(e)}
    
    world_bank = {
        "source": "World Bank Commodity Markets - Pink Sheet",
        "release_date": "2026-04-03",
        "confidence": "HIGH",
        "url": "https://www.worldbank.org/en/research/commodity-markets",
        "march_2026_changes": {
            "fertilizer_index_mom_pct": 26.2,
            "urea_mom_pct": 53.7,
            "energy_index_mom_pct": 41.6,
            "natural_gas_europe_mom_pct": 59.4,
            "crude_oil_mom_pct_range": "40.5-45.8"
        },
        "raw_data": wb_data_raw
    }

    # ==========================================
    # 4. FUEL DATA (April 2026)
    # ==========================================
    fuel_2026 = {
        "diesel_retail_clear": {
            "price_per_litre": 2.37,
            "source": "CTV News, early April 2026",
            "confidence": "MEDIUM",
            "notes": "Retail pump price, not farm bulk"
        },
        "diesel_farm_bulk_estimated": {
            "conservative": 2.22,
            "mid": 2.18,
            "optimistic": 2.04,
            "working_estimate": 2.15,
            "unit": "CAD/litre",
            "methodology": "Applied validated 14% / 19c/L discount from Ridgetown 2025 comparison",
            "tax_components": {
                "provincial_fuel_tax_exemption": 0.09,
                "bulk_delivery_margin_savings": 0.055,
                "hst_on_tax_diff": 0.012,
                "volume_pricing": 0.035
            }
        },
        "brent_crude": {
            "price_per_barrel_usd": 110.0,
            "date": "2026-04-07",
            "source": "Trading Economics / multiple outlets"
        },
        "propane": {
            "lock_in_2025": {"low": 0.69, "high": 0.75, "unit": "CAD/litre"},
            "spot_2026_estimated": {"low": 0.95, "high": 1.15, "unit": "CAD/litre",
                                    "notes": "Estimated based on crude oil price correlation"}
        }
    }

    # ==========================================
    # 5. EXCHANGE RATE
    # ==========================================
    exchange_rate = {
        "cad_usd": 0.7185,
        "usd_cad": 1.3919,
        "date": "2026-04-07",
        "source": "XE/MTFX"
    }

    # ==========================================
    # 6. ESTIMATED APRIL 2026 ONTARIO PRICES
    # ==========================================
    
    # Method A: World Bank escalation
    urea_method_a = round(819.0 * (1 + 0.537), 0)
    # Method B: DTN cross-reference with ON discount
    urea_method_b = round(826.0 * 1.1023 * (1 / 0.7185) * 0.94, 0)
    # Method C: News midpoint (45%)
    urea_method_c = round(819.0 * 1.45, 0)
    
    # MAP: Use fertilizer index (+26.2%) as proxy
    map_method_a = round(1164.0 * (1 + 0.262), 0)
    map_method_b = round(857.0 * 1.1023 * (1 / 0.7185) * 0.94, 0)  # Using DAP as MAP proxy
    
    # Potash: Use fertilizer index (+26.2%) as proxy
    potash_method_a = round(678.0 * (1 + 0.262), 0)
    potash_method_b = round(489.0 * 1.1023 * (1 / 0.7185) * 0.94, 0)
    
    estimated_2026 = {
        "methodology": "Three-way triangulation: World Bank index escalation (primary), DTN cross-reference, news validation",
        "confidence": "MEDIUM - derived estimates, not direct Ontario survey data",
        "fertilizer": {
            "urea_46_0_0": {
                "method_a_wb_escalation": urea_method_a,
                "method_b_dtn_crossref": urea_method_b,
                "method_c_news_midpoint": urea_method_c,
                "working_range": {"low": min(urea_method_b, urea_method_c), 
                                  "high": urea_method_a},
                "midpoint": round((urea_method_a + urea_method_b + urea_method_c) / 3, 0),
                "unit": "CAD/metric tonne",
                "yoy_change_pct_range": {
                    "low": round((min(urea_method_b, urea_method_c) / 819.0 - 1) * 100, 1),
                    "high": round((urea_method_a / 819.0 - 1) * 100, 1)
                }
            },
            "map_11_52_0": {
                "method_a_wb_escalation": map_method_a,
                "method_b_dtn_crossref": map_method_b,
                "working_range": {"low": min(map_method_a, map_method_b),
                                  "high": max(map_method_a, map_method_b)},
                "midpoint": round((map_method_a + map_method_b) / 2, 0),
                "unit": "CAD/metric tonne"
            },
            "potash_mop_0_0_60": {
                "method_a_wb_escalation": potash_method_a,
                "method_b_dtn_crossref": potash_method_b,
                "working_range": {"low": min(potash_method_a, potash_method_b),
                                  "high": max(potash_method_a, potash_method_b)},
                "midpoint": round((potash_method_a + potash_method_b) / 2, 0),
                "unit": "CAD/metric tonne"
            }
        },
        "fuel": {
            "diesel_coloured_farm": {
                "working_range": {"low": 2.04, "high": 2.22},
                "midpoint": 2.15,
                "unit": "CAD/litre",
                "yoy_change_pct": round((2.15 / 1.16 - 1) * 100, 1)
            }
        }
    }

    # ==========================================
    # 7. APPLICATION RATES (OMAFRA-based)
    # ==========================================
    application_rates = {
        "source": "OMAFRA Publications 60, 811, 839 + Ontario Corn N Calculator",
        "grain_crops": {
            "corn": {
                "n_rate_lbs_per_acre": {"typical": 150, "range": [130, 170]},
                "p2o5_rate_lbs_per_acre": {"typical": 50, "range": [30, 70]},
                "k2o_rate_lbs_per_acre": {"typical": 80, "range": [50, 100]},
                "diesel_litres_per_acre": {"min_till": 25, "conventional": 40},
                "yield_bu_per_acre": {"typical": 180, "range": [160, 200]},
                "drying_moisture_removal_points": 10,
                "propane_gal_per_point_per_bu": 0.02,
                "notes": "N rate highly variable — use Ontario Corn N Calculator for precision"
            },
            "soybeans": {
                "n_rate_lbs_per_acre": 0,
                "p2o5_rate_lbs_per_acre": {"typical": 30, "range": [20, 50]},
                "k2o_rate_lbs_per_acre": {"typical": 60, "range": [40, 80]},
                "diesel_litres_per_acre": {"min_till": 20, "conventional": 35},
                "notes": "Soybeans fix their own nitrogen"
            },
            "winter_wheat": {
                "n_rate_lbs_per_acre": {"typical": 70, "range": [50, 90]},
                "p2o5_rate_lbs_per_acre": {"typical": 40, "range": [25, 55]},
                "k2o_rate_lbs_per_acre": {"typical": 50, "range": [30, 70]},
                "diesel_litres_per_acre": {"min_till": 20, "conventional": 35}
            }
        },
        "vegetable_crops": {
            "source": "OMAFRA Publication 839 - Guide to Vegetable Production",
            "sweet_corn": {"n_rate_lbs_per_acre": {"typical": 135, "range": [120, 150]}},
            "potatoes": {"n_rate_lbs_per_acre": {"typical": 175, "range": [150, 200]}},
            "tomatoes_processing": {"n_rate_lbs_per_acre": {"typical": 120, "range": [100, 140]}},
            "peppers": {"n_rate_lbs_per_acre": {"typical": 115, "range": [100, 130]}},
            "leafy_greens": {"n_rate_lbs_per_acre": {"typical": 100, "range": [80, 120]}}
        },
        "conversions": {
            "lbs_per_tonne": 2204.6,
            "lbs_per_us_ton": 2000,
            "litres_per_us_gallon": 3.785,
            "urea_n_content": 0.46,
            "map_n_content": 0.11,
            "map_p2o5_content": 0.52,
            "dap_n_content": 0.18,
            "dap_p2o5_content": 0.46,
            "mop_k2o_content": 0.60,
            "uan28_n_content": 0.28
        }
    }

    # ==========================================
    # 8. GEOPOLITICAL CONTEXT
    # ==========================================
    context = {
        "conflict": "Iran War / Strait of Hormuz Disruption",
        "onset": "Late February 2026",
        "key_facts": [
            "Strait of Hormuz traffic declined >90%",
            "~20-30% of global oil transits Hormuz",
            "~1/3 of global fertilizer trade transits Hormuz",
            "Brent crude surged to $110+/barrel",
            "Natural gas (Europe TTF) up ~59% in March 2026",
            "Urea prices globally up 40-54% since conflict",
            "Ontario farmers entered spring with low on-farm fertilizer inventory",
            "Federal carbon tax zeroed April 1, 2025 (no longer a factor)",
            "Canada 35% tariff on Russian fertilizer still in effect"
        ],
        "ontario_vulnerability": {
            "potash": "LOW - Canada is world's largest producer (Saskatchewan)",
            "nitrogen_urea": "HIGH - Eastern Canada (incl. Ontario) is net importer of finished N products",
            "phosphate_dap_map": "HIGH - Canada has very limited domestic phosphate production"
        }
    }

    # ==========================================
    # ASSEMBLE COMPLETE DATASET
    # ==========================================
    dataset = {
        "metadata": {
            "title": "Ontario Farm Energy & Fertilizer Cost Impact Dataset",
            "purpose": "Year-over-year comparison: April 2025 vs April 2026",
            "generated_at": datetime.now().isoformat(),
            "generated_by": "input_cost_data_collector.py",
            "version": "1.0"
        },
        "baseline_2025": baseline_2025,
        "dtn_march_2026": dtn_march_2026,
        "world_bank": world_bank,
        "exchange_rate": exchange_rate,
        "fuel_april_2026": fuel_2026,
        "estimated_april_2026_ontario": estimated_2026,
        "application_rates": application_rates,
        "geopolitical_context": context
    }

    return dataset


def main():
    print("=" * 60)
    print("Ontario Farm Cost Impact Data Collector")
    print("=" * 60)
    
    dataset = build_dataset()
    
    # Save to JSON
    output_path = os.path.join(os.path.dirname(os.path.dirname(__file__)), 
                               "data", "farm_cost_impact_data.json")
    with open(output_path, "w", encoding="utf-8") as f:
        json.dump(dataset, f, indent=2, ensure_ascii=False)
    
    print(f"\n✅ Dataset saved to: {output_path}")
    
    # Print summary
    est = dataset["estimated_april_2026_ontario"]["fertilizer"]
    base = dataset["baseline_2025"]["fertilizer"]
    fuel_base = dataset["baseline_2025"]["fuel"]
    fuel_est = dataset["estimated_april_2026_ontario"]["fuel"]
    
    print("\n" + "=" * 60)
    print("PRICE COMPARISON SUMMARY")
    print("=" * 60)
    
    print(f"\n{'Product':<20} {'2025 Baseline':>15} {'2026 Range':>20} {'2026 Midpoint':>15} {'YoY Change':>12}")
    print("-" * 85)
    
    for key in ["urea_46_0_0", "map_11_52_0", "potash_mop_0_0_60"]:
        b = base[key]["price_per_tonne"]
        e = est[key]
        lo = e["working_range"]["low"]
        hi = e["working_range"]["high"]
        mid = e["midpoint"]
        pct = round((mid / b - 1) * 100, 1)
        name = key.split("_")[0].upper()
        print(f"{name:<20} ${b:>13,.0f}/t ${lo:>7,.0f}-${hi:>5,.0f}/t ${mid:>13,.0f}/t {pct:>+10.1f}%")
    
    # Fuel
    fb = fuel_base["diesel_coloured_farm"]["price_per_litre"]
    fe = fuel_est["diesel_coloured_farm"]
    fmid = fe["midpoint"]
    fpct = round((fmid / fb - 1) * 100, 1)
    print(f"{'Farm Diesel':<20} ${fb:>13.2f}/L ${fe['working_range']['low']:>7.2f}-${fe['working_range']['high']:>5.2f}/L ${fmid:>13.2f}/L {fpct:>+10.1f}%")
    
    print("\n" + "=" * 60)
    print("Data collection complete. Ready for modelling.")
    print("=" * 60)


if __name__ == "__main__":
    main()
