"""
Inspect FIR Data Structure
==========================
Reads a sample FIR xlsx file and dumps all sheet names, headers, and sample rows
to a UTF-8 report file for analysis.
"""

import openpyxl
from pathlib import Path

SAMPLE_FILE = Path(r"C:\Projects\Farm Finance Stats Dashboard\Database\farm_finance_dashboard_starter\data\FIR Data\_fir_inspect_temp\FI240101 South Glengarry Tp.xlsx")
REPORT_FILE = Path(r"C:\Projects\Farm Finance Stats Dashboard\Database\farm_finance_dashboard_starter\data\FIR Data\_fir_inspect_temp\fir_inspection_report.md")


def inspect():
    wb = openpyxl.load_workbook(str(SAMPLE_FILE), data_only=True)
    
    lines = []
    lines.append("# FIR Data Inspection Report")
    lines.append(f"**Source**: {SAMPLE_FILE.name}")
    lines.append(f"**Total Sheets**: {len(wb.sheetnames)}")
    lines.append("")
    
    # Sheet summary table
    lines.append("## Sheet Summary")
    lines.append("")
    lines.append("| # | Sheet | Rows | Cols |")
    lines.append("|---|-------|------|------|")
    for i, name in enumerate(wb.sheetnames, 1):
        ws = wb[name]
        lines.append(f"| {i} | {name} | {ws.max_row} | {ws.max_column} |")
    lines.append("")
    
    # Detailed per-sheet inspection
    for name in wb.sheetnames:
        ws = wb[name]
        lines.append(f"---")
        lines.append(f"## Sheet: {name}")
        lines.append(f"**Dimensions**: {ws.max_row} rows x {ws.max_column} cols")
        lines.append("")
        
        # Get all rows as values
        all_rows = []
        for row in ws.iter_rows(min_row=1, max_row=min(ws.max_row, 20), values_only=False):
            row_data = []
            for cell in row:
                row_data.append(str(cell.value) if cell.value is not None else "")
            all_rows.append(row_data)
        
        # Print first 15 rows to understand the structure
        if all_rows:
            lines.append("**First rows (up to 15):**")
            lines.append("```")
            for r_idx, row in enumerate(all_rows[:15]):
                # Truncate long values
                formatted = [v[:60] if len(v) > 60 else v for v in row]
                lines.append(f"  Row {r_idx+1}: {formatted}")
            lines.append("```")
            lines.append("")
        
        # Also show rows from the middle of the sheet
        if ws.max_row > 20:
            mid = ws.max_row // 2
            lines.append(f"**Sample rows from middle (rows {mid}-{mid+5}):**")
            lines.append("```")
            for row in ws.iter_rows(min_row=mid, max_row=min(mid+5, ws.max_row), values_only=False):
                row_data = [str(c.value)[:60] if c.value is not None else "" for c in row]
                lines.append(f"  Row {row[0].row}: {row_data}")
            lines.append("```")
            lines.append("")
    
    report = "\n".join(lines)
    REPORT_FILE.write_text(report, encoding="utf-8")
    print(f"Report written to: {REPORT_FILE}")
    print(f"Size: {len(report)} characters")


if __name__ == "__main__":
    inspect()
