"""Extract full report template structure."""
from docx import Document

doc = Document(r'C:\Users\ben.lefort\OneDrive - Ontario Federation of Agriculture\Desktop\Ben Desktop Files\Economic Analyst Position\Economic Impact Studies\County Level\Perth and Wellington\Report.docx')

print("=== FULL DOCUMENT STRUCTURE ===")
print()
for i, p in enumerate(doc.paragraphs):
    if p.text.strip():
        print(f"[{i:3d}] Style='{p.style.name}' | Text='{p.text}'")

print("\n=== TABLES ===")
for i, table in enumerate(doc.tables):
    print(f"\nTable {i}: {len(table.rows)} rows x {len(table.columns)} cols")
    for j, row in enumerate(table.rows):
        cells = [cell.text.strip() for cell in row.cells]
        print(f"  Row {j}: {cells}")
