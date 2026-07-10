"""
Direct edit of the existing Word document to:
1. Update urea speaking notes with Trading Economics live source citation
2. Update methodology notes with Trading Economics reference
"""

import sys
from docx import Document

DOC_PATH = (
    r"C:\Users\ben.lefort\OneDrive - Ontario Federation of Agriculture"
    r"\Desktop\Ben Desktop Files\Farm Business Confidence Survey"
    r"\2025\Specififc Issue spekaing notes"
    r"\Energy and Fertilizer Shock Memo and Speaking Notes.docx"
)

# Load existing document
doc = Document(DOC_PATH)

# Track what we changed
changes = []

# Search and replace in all paragraphs
for i, para in enumerate(doc.paragraphs):
    text = para.text

    # --- FIX 1: Update the urea speaking note in grain farm talking points ---
    if ("Urea fertilizer" in text and "$819 per tonne" in text 
            and "approximately $1,200" in text):
        # Preserve the formatting of the first run, replace text
        new_text = (
            "Urea fertilizer, the primary source of nitrogen for corn and wheat, "
            "cost $819 per tonne in Ontario last spring. Global urea benchmark "
            "prices are currently trading above $700 USD per tonne (Trading Economics, "
            "April 7, 2026). After accounting for freight, retail margins, and currency "
            "conversion, Ontario farm-gate urea is estimated at approximately $1,200 "
            "per tonne \u2014 a 45 to 54 percent increase driven directly by the energy "
            "crisis in the Middle East."
        )
        # Clear existing runs and set new text while preserving paragraph style
        if para.runs:
            # Keep formatting from first run
            fmt = para.runs[0].font
            style_name = para.style.name
            for run in para.runs:
                run.text = ""
            para.runs[0].text = new_text
        else:
            para.text = new_text
        changes.append(f"Paragraph {i}: Updated urea talking point with Trading Economics citation")

    # --- FIX 2: Update the urea line in the "estimates conservative or aggressive" Q&A ---
    # Look for the Q&A answer about estimates
    if ("average of three independent estimation methods" in text 
            and "World Bank" in text and "DTN" in text):
        new_text = (
            "A: We present three scenarios. Our conservative estimate uses the lowest "
            "available price data; our worst-case uses the highest. Our midpoint is the "
            "average of three independent estimation methods, anchored to publicly "
            "verifiable data: current global urea benchmark prices (Trading Economics), "
            "the World Bank\u2019s March 2026 commodity index, and the DTN U.S. agricultural "
            "retail survey. When Statistics Canada releases its Farm Input "
            "Price Index for Q4 2025 on April 10, we will incorporate that as additional "
            "validation."
        )
        if para.runs:
            for run in para.runs:
                run.text = ""
            para.runs[0].text = new_text
        else:
            para.text = new_text
        changes.append(f"Paragraph {i}: Updated Q&A answer with Trading Economics reference")

# Also search within tables
for table_idx, table in enumerate(doc.tables):
    for row_idx, row in enumerate(table.rows):
        for cell_idx, cell in enumerate(row.cells):
            for para in cell.paragraphs:
                # Update source citation row if it exists
                if "2026 estimates triangulated" in para.text:
                    new_source = (
                        "2025 baselines from Ridgetown Farm Input Monitoring Project "
                        "(April 16, 2025). 2026 estimates triangulated from World Bank "
                        "Pink Sheet (March 2026), DTN Retail Fertilizer Index (March 2026), "
                        "and validated against Trading Economics global urea benchmark "
                        "($701 USD/tonne, April 7, 2026). Farm diesel estimated using "
                        "validated 14% farm-bulk discount on retail prices."
                    )
                    if para.runs:
                        for run in para.runs:
                            run.text = ""
                        para.runs[0].text = new_source
                    else:
                        para.text = new_source
                    changes.append(f"Table {table_idx}, Row {row_idx}: Updated source citation with Trading Economics")

# Save
doc.save(DOC_PATH)

print(f"Document saved to: {DOC_PATH}")
print(f"\nChanges made ({len(changes)}):")
for c in changes:
    print(f"  - {c}")

if not changes:
    print("\n  WARNING: No matching text found. The document structure may have changed.")
    print("  Dumping first 50 paragraphs for inspection:")
    doc2 = Document(DOC_PATH)
    for i, p in enumerate(doc2.paragraphs[:50]):
        if p.text.strip():
            print(f"    [{i}] {p.text[:120]}...")
